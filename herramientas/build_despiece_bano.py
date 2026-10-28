"""Genera despiece-bano.html desde los valores guardados en despiece-bano.xlsx.
Uso: python herramientas/build_despiece_bano.py
"""
from pathlib import Path
import html, re, sys
import openpyxl
ROOT=Path(__file__).resolve().parent.parent

def text(v):
 if v is None:return '—'
 if isinstance(v,(int,float)):return f'{v:,.6f}'.rstrip('0').rstrip('.').replace(',','@').replace('.',',').replace('@','.')
 value=html.escape(str(v))
 return re.sub(r'https?://[^\s;()<>]+',lambda m:'<a target="_blank" rel="noopener" href="'+m.group(0)+'">Fuente ↗</a>',value)

def build():
 source=ROOT/'despiece-bano.xlsx'
 book=openpyxl.load_workbook(source,data_only=True)
 formulas=openpyxl.load_workbook(source,data_only=False)
 for ws in formulas:
  for row in ws:
   for cell in row:
    if cell.data_type=='f' and book[ws.title][cell.coordinate].value is None:raise ValueError(f'Falta valor calculado: {ws.title}!{cell.coordinate}')
 sections=[]
 for ws in book:
  cols=[i for i in range(1,ws.max_column+1) if ws.cell(6,i).value is not None]
  lines=[]
  for row in range(7,ws.max_row+1):
   if not any(ws.cell(row,i).value is not None for i in cols):continue
   cells=[]
   for col in cols:
    val=ws.cell(row,col).value
    if ws.title=='Compras' and col in (5,6) and isinstance(val,(int,float)):v='$'+f'{val:,.0f}'.replace(',','.')
    else:v=text(val)
    cells.append(f'<td class="num">{v}</td>' if isinstance(val,(int,float)) else f'<td>{v}</td>')
   lines.append('<tr>'+''.join(cells)+'</tr>')
  headers=''.join('<th>'+text(ws.cell(6,i).value)+'</th>' for i in cols)
  intro=''.join('<p class="note">'+text(ws.cell(row,1).value)+'</p>' for row in (3,4) if ws.cell(row,1).value)
  sections.append(f'<section id="{html.escape(ws.title)}"><h2>{text(ws.cell(2,1).value)}</h2>{intro}<div class="table"><table><thead><tr>{headers}</tr></thead><tbody>'+''.join(lines)+'</tbody></table></div></section>')
 notice='Mueble flotante y gabinete con espejo. Cantos completos por zona húmeda. Tablero, fondo, canto y bisagras se compran con cocina y clóset en compras-proyecto.xlsx; antes de cortar, revisar la hoja Pendientes. Precios parciales.'
 doc='''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Despiece del conjunto de baño</title><style>
 :root{color-scheme:light}*{box-sizing:border-box}body{margin:0;background:#f7f7f0;color:#253b35;font:15px/1.6 system-ui,sans-serif}main{max-width:1320px;margin:auto;padding:24px 18px 48px}h1{font:34px Georgia,serif}h2{font:25px Georgia,serif;margin-top:36px}a{color:#355b46}nav{display:flex;gap:16px;flex-wrap:wrap;padding:12px 0}.notice{padding:14px 18px;background:#fff0cb;border-left:4px solid #b78a36;border-radius:6px}.note{color:#59645d;margin:8px 0}.table{overflow:auto;border:1px solid #d5dbce;border-radius:8px;background:white}table{width:100%;border-collapse:collapse;font-size:13px}th{background:#253b35;color:white;text-align:left;white-space:nowrap}td,th{padding:12px;border-bottom:1px solid #e6e9df;vertical-align:top}td{min-width:100px;max-width:530px}td:last-child{min-width:300px}.num{text-align:right;white-space:nowrap;min-width:76px}tr:nth-child(even){background:#f5f7f0}footer{margin-top:24px;font-size:13px;color:#677165}@media(max-width:600px){main{padding:18px 12px}h1{font-size:28px}}
 </style></head><body><main><a href="index.html">← Espacios 3D</a><h1>Baño · despiece y compras</h1><nav><a href="bano.html">Abrir visor 3D ↗</a><a download href="'''+source.name+'''">Descargar Excel ↓</a>'''+''.join('<a href="#'+html.escape(ws.title)+'">'+html.escape(ws.title)+'</a>' for ws in book)+'''</nav><p class="notice">'''+html.escape(notice)+'''</p>'''+''.join(sections)+'''<footer>Vista de solo lectura, generada desde el Excel. Regenerar esta página después de modificar y guardar sus cálculos.</footer></main></body></html>'''
 dest=ROOT/'despiece-bano.html';dest.write_text(doc,encoding='utf-8');print(f'OK: {dest.name}')
if __name__=='__main__':build()
