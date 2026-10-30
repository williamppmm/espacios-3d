"""Genera el presupuesto y el pedido de corte conjuntos de los tres muebles.

Uso, desde la raíz del repositorio:

    python herramientas/build_compras_proyecto.py

Lee despiece-cocina.xlsx, despiece-closet.xlsx y despiece-bano.xlsx (valores guardados por Excel)
y los manifiestos assets/models/*-data.js. Escribe compras-proyecto.xlsx y compras-proyecto.html.
Los Excel de cada mueble siguen siendo la fuente: si cambia un despiece, guárdalo en Excel y vuelve
a ejecutar este script. Requisitos: Python 3.9+ y openpyxl.
"""
import html, json, math, random, sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from build_despiece import page, tabla, fmt

RAIZ = Path(__file__).resolve().parent.parent
FECHA = "29/10/2026"
PRECIO = {"RH15": 365872, "RH6": 212534, "rigido": 531600, "flexible": 285600, "bisagra_par": 17750}
ROLLO = {"rigido": 100, "flexible": 200}
RESERVA_BORDE, MARGEN = 0.05, 0.10   # m por borde y margen sobre el total de canto

MUEBLES = [
    # nombre, Excel, manifiesto, Despiece: fila inicial, columnas de medida terminada y de corte (largo, ancho),
    # celdas de Cantos (neto, bordes)
    dict(nombre="Cocina", xlsx="despiece-cocina.xlsx", js="cocina-data.js", fila=5, fin=(5, 6), corte=(12, 13),
         cantos={"2": ("C4", "D4"), "0.45": ("C5", "D5")}),
    dict(nombre="Clóset", xlsx="despiece-closet.xlsx", js="closet-data.js", fila=7, fin=(5, 6), corte=(13, 14),
         cantos={"2": ("C7", "D7"), "0.45": ("C8", "D8")}),
    dict(nombre="Baño", xlsx="despiece-bano.xlsx", js="bano-data.js", fila=7, fin=(6, 7), corte=(14, 15),
         cantos={"2": ("B7", "C7"), "0.45": ("B8", "C8")}),
]


def libro(nombre):
    return openpyxl.load_workbook(RAIZ / nombre, data_only=True)


def num(v):
    return v if isinstance(v, (int, float)) else None


# ------------------------------------------------------------ pedido de corte conjunto
NO_PEDIR = "NO PEDIR"   # piezas que se cortan en el armado, de un sobrante
EN_ARMADO = ("NO PEDIR", "Rebaje", "recortar", "Pedida", "corte central", "Recorte")


def pedido():
    """Filas del pedido de cada mueble y sobremedida pedida (largo y ancho a pedir menos la medida terminada)."""
    filas, sobremedida = [], {}
    for m in MUEBLES:
        wb = libro(m["xlsx"])
        ds, pc = wb["Despiece"], wb["Pedido de corte"]
        cortes = {}
        for r in range(m["fila"], ds.max_row + 1):
            pid = ds.cell(r, 1).value
            if isinstance(pid, str) and ds.cell(r, m["corte"][0]).value is not None:
                cortes[pid] = (ds.cell(r, m["corte"][0]).value, ds.cell(r, m["corte"][1]).value,
                               ds.cell(r, m["fin"][0]).value, ds.cell(r, m["fin"][1]).value)
        for r in range(7, pc.max_row + 1):
            v = [pc.cell(r, c).value for c in range(1, 15)]
            if not v[0] or v[0] not in cortes:
                continue
            if {"R", "F"}.issubset(set(v[7:11])):
                raise ValueError(f"{m['nombre']} {v[0]}: el optimizador admite un solo tipo de canto por pieza")
            cl, ca, fl, fa = cortes[v[0]]
            extra = lambda pedir, fin: pedir - fin if num(pedir) is not None and num(fin) is not None and pedir > fin else 0
            dl, da = extra(v[5], fl), extra(v[6], fa)
            sobremedida[f'{m["nombre"]} {v[0]}'] = (dl, da)
            mm = lambda x: f"{round(x, 2):g}".replace(".", ",")   # sin separador de miles: 2171, no 2.171
            filas.append([m["nombre"]] + v[:13] + [f"{mm(cl + dl)} × {mm(ca + da)}", v[13] or "—"])
    return filas, sobremedida


# ------------------------------------------------------------ encaje conjunto en láminas
K, BW, BH = 5, 2440 - 20, 2150 - 20   # sierra 5 mm según PDF Madecentro; refilado 10 mm por lado


def piezas_de_manifiestos(excluir, sobremedida):
    canto = lambda e: 2 if e == "2" else 0.45 if e == "0.45" else 0
    rh, fondo = [], []
    for m in MUEBLES:
        t = (RAIZ / "assets" / "models" / m["js"]).read_text(encoding="utf-8")
        d = json.loads(t[t.index("{"): t.rindex("}") + 1])
        for p in d["parts"]:
            mat = p.get("material", "RH15")
            esp = p.get("t") or (6 if " 6" in mat or "RH6" in mat else 15)
            clave = f'{m["nombre"]} {p["id"]}'
            if clave in excluir:
                continue
            dl, da = sobremedida.get(clave, (0, 0))
            w = p["L"] - canto(p["edges"][2]) - canto(p["edges"][3]) + dl
            h = p["A"] - canto(p["edges"][0]) - canto(p["edges"][1]) + da
            for _ in range(p["qty"]):
                (fondo if esp == 6 else rh).append((clave, w, h))
    return rh, fondo


def empacar(orden, elegir, dividir):
    laminas = []
    for pid, w, h in orden:
        W, H = w + K, h + K
        for lam in laminas + [None]:
            if lam is None:
                lam = {"libre": [(0, 0, BW + K, BH + K)], "items": []}
                laminas.append(lam)
            mejor = None
            for i, (x, y, fw, fh) in enumerate(lam["libre"]):
                if W <= fw and H <= fh:
                    s = {"baf": fw * fh - W * H, "bssf": min(fw - W, fh - H), "blsf": max(fw - W, fh - H)}[elegir]
                    if mejor is None or s < mejor[0]:
                        mejor = (s, i)
            if mejor:
                x, y, fw, fh = lam["libre"].pop(mejor[1])
                lam["items"].append((pid, x, y, w, h))
                rw, rh = fw - W, fh - H
                if {"sas": rw < rh, "las": rw >= rh}[dividir]:
                    if rw > 0: lam["libre"].append((x + W, y, rw, H))
                    if rh > 0: lam["libre"].append((x, y + H, fw, rh))
                else:
                    if rw > 0: lam["libre"].append((x + W, y, rw, fh))
                    if rh > 0: lam["libre"].append((x, y + H, W, rh))
                break
            if lam is laminas[-1] and not lam["items"]:
                return None
    return laminas


def mejor_encaje(piezas, intentos=600):
    rnd = random.Random(7)
    ordenes = [sorted(piezas, key=lambda p: -p[1] * p[2]), sorted(piezas, key=lambda p: -p[1]), sorted(piezas, key=lambda p: -p[2])]
    for _ in range(intentos):
        o = list(piezas); rnd.shuffle(o); ordenes.append(o)
    mejor = None
    for o in ordenes:
        for e in ("baf", "bssf", "blsf"):
            for s in ("sas", "las"):
                r = empacar(o, e, s)
                if r and (mejor is None or len(r) < len(mejor)):
                    mejor = r
    return mejor


def ocupacion(lam):
    return sum(w * h for _, _, _, w, h in lam["items"]) / (BW * BH)


# ------------------------------------------------------------ cantos
def cantos():
    res = {}
    for tipo in ("2", "0.45"):
        por_mueble = []
        for m in MUEBLES:
            ws = libro(m["xlsx"])["Cantos"]
            neto, bordes = (ws[c].value for c in m["cantos"][tipo])
            por_mueble.append((m["nombre"], float(neto), int(bordes)))
        neto = sum(x[1] for x in por_mueble); bordes = sum(x[2] for x in por_mueble)
        necesario = math.ceil((neto + bordes * RESERVA_BORDE) * (1 + MARGEN))
        res[tipo] = dict(por_mueble=por_mueble, neto=neto, bordes=bordes, necesario=necesario)
    r = res["2"]; r["rollos"] = 1; r["sueltos"] = max(0, r["necesario"] - ROLLO["rigido"])
    f = res["0.45"]; f["rollos"] = math.ceil(f["necesario"] / ROLLO["flexible"]); f["sueltos"] = 0
    f["sobrante"] = f["rollos"] * ROLLO["flexible"] - f["necesario"]
    return res


# ------------------------------------------------------------ compras consolidadas
GENERICOS = ("Por elegir", "Por dimensionar", "Por cotizar", "Pino · por cotizar", "Verificar con FSCPCM20")


def compras(n_rh, n_fondo, canto):
    filas = []   # grupo, ítem, mueble, cantidad, unidad, precio, nota
    add = lambda *f: filas.append(list(f))
    add("Tableros y cantos", "Tablero melamina Cartagena RH15 · 2440 × 2150", "Cocina, clóset y baño", n_rh, "lámina", PRECIO["RH15"],
        f"Encaje conjunto de las piezas RH15 de los tres muebles en {n_rh} láminas (hoja Láminas). El total por separado se muestra en Resumen.")
    add("Tableros y cantos", "Fondo melamina blanco RH6 · 2440 × 2150", "Cocina y baño", n_fondo, "lámina", PRECIO["RH6"],
        "Los 9 fondos de la cocina y el fondo del gabinete del baño caben en una lámina.")
    r, f = canto["2"], canto["0.45"]
    add("Tableros y cantos", "Canto rígido Cartagena 19 × 2 mm · rollo 100 m", "Cocina, clóset y baño", r["rollos"], "rollo", PRECIO["rigido"],
        f"Necesario {r['necesario']} m: {fmt(round(r['neto'], 1))} m netos + 5 cm por borde ({r['bordes']} bordes) + 10 %.")
    add("Tableros y cantos", "Canto rígido Cartagena 19 × 2 mm · metros sueltos", "Cocina, clóset y baño", r["sueltos"], "m", "cotizar",
        "Complemento del rollo de 100 m; se compra suelto en la sucursal.")
    add("Tableros y cantos", "Canto flexible Cartagena 19 × 0,45 mm · rollo 200 m", "Cocina, clóset y baño", f["rollos"], "rollo", PRECIO["flexible"],
        f"Necesario {f['necesario']} m: {fmt(round(f['neto'], 1))} m netos + 5 cm por borde ({f['bordes']} bordes) + 10 %. "
        f"Sobran unos {f['sobrante']} m para cantear en el armado las piezas pequeñas que no acepte la máquina.")
    # Cocina
    co = libro("despiece-cocina.xlsx")["Compras"]
    bisagras_cocina = 0
    for r_ in range(2, co.max_row + 1):
        g, item, cant, precio, _, nota = (co.cell(r_, c).value for c in range(1, 7))
        if not g or g == "Tableros y cantos":
            continue
        if str(item).startswith("Bisagra"):
            bisagras_cocina = cant; continue
        if str(item).startswith("Corte y canteado"):
            continue
        grupo = "Mesón, equipos y grifería" if g == "Mesón y equipos" else g
        add(grupo, item, "Cocina", cant, "", precio, nota)
    # Clóset
    cl = libro("despiece-closet.xlsx")["Compras"]
    for r_ in range(10, cl.max_row + 1):
        ref, uso, cant, unidad, precio, _, _, det = (cl.cell(r_, c).value for c in range(1, 9))
        if not ref or str(ref).startswith("SUBTOTAL"):
            continue
        if str(uso).startswith("Corte, canteado"):
            continue
        item = uso if ref in GENERICOS else f"{ref} · {uso}"
        add("Herrajes", item, "Clóset", cant, unidad, precio, det)
    # Baño
    ba = libro("despiece-bano.xlsx")["Compras"]
    bisagras_bano = 0
    for r_ in range(7, ba.max_row + 1):
        ref, uso, cant, unidad, precio, _, _, det = (ba.cell(r_, c).value for c in range(1, 9))
        if not ref or str(ref).startswith("REFERENCIAS"):
            continue
        if ref.startswith(("RH15", "Fondo", "Canto", "Corte")):
            continue
        if ref.startswith("HBM897"):
            bisagras_bano = cant; continue
        grupo = ("Otros proveedores" if ref in ("Espejo", "Retal de granito") else
                 "Mesón, equipos y grifería" if ref in ("Lavamanos de sobreponer", "Grifería", "Desagüe y sifón") else "Herrajes")
        add(grupo, f"{ref} · {uso}", "Baño", cant, unidad, precio, det)
    # Bisagras compartidas: misma referencia en cocina y baño.
    add("Herrajes", "Bisagra parche HBM897-01 · par", "Cocina y baño", (bisagras_cocina or 0) + (bisagras_bano or 0), "par", PRECIO["bisagra_par"],
        f"Cocina {bisagras_cocina} pares (24 bisagras) + baño {bisagras_bano} pares (6 bisagras). Validar peso y cabida en el gabinete del espejo.")
    add("Servicios", "Corte y canteado · pedido único", "Cocina, clóset y baño", 1, "servicio", "cotizar",
        "Un solo pedido a la sucursal (hoja Pedido de corte). No cortar las piezas marcadas en Bloqueos hasta resolverlas.")
    orden = ["Tableros y cantos", "Herrajes", "Mesón, equipos y grifería", "Otros proveedores", "Accesorios", "Servicios"]
    filas.sort(key=lambda f: orden.index(f[0]))
    return filas, orden


def subtotal(f):
    return f[3] * f[5] if isinstance(f[3], (int, float)) and isinstance(f[5], (int, float)) else None


# ------------------------------------------------------------ escritura
TINTA = PatternFill("solid", fgColor="253B35")


def encabezado(ws, fila, titulos, anchos):
    for i, t in enumerate(titulos, 1):
        c = ws.cell(fila, i, t); c.font = Font(bold=True, color="FFFFFF"); c.fill = TINTA
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = a


def titulo(ws, t, nota):
    ws["A2"] = t; ws["A2"].font = Font(bold=True, size=14)
    ws["A3"] = nota; ws["A3"].font = Font(italic=True)


def main():
    todas, sobremedida = pedido()
    filas_pedido = [f for f in todas if NO_PEDIR not in str(f[15])]
    armado = [f for f in todas if any(k in str(f[15]) for k in EN_ARMADO)]
    rh, fondo = piezas_de_manifiestos({f"{f[0]} {f[1]}" for f in todas if NO_PEDIR in str(f[15])}, sobremedida)
    lam_rh, lam_fondo = mejor_encaje(rh), mejor_encaje(fondo)
    canto = cantos()
    filas_compras, grupos = compras(len(lam_rh), len(lam_fondo), canto)
    bloqueos = [f for f in filas_pedido if any(k in str(f[15]) for k in ("NO CORTAR", "antes de cortar", "Cortar después"))]
    pend_cl = libro("despiece-closet.xlsx")["Pendientes"]
    for r in range(7, pend_cl.max_row + 1):
        if pend_cl.cell(r, 4).value == "Todo el corte":
            bloqueos.insert(0, ["Clóset", "Todo", "", pend_cl.cell(r, 2).value] + [""] * 11 + [pend_cl.cell(r, 3).value])
    pend_co = libro("despiece-cocina.xlsx")["Pendientes"]
    for r in range(7, pend_co.max_row + 1):
        if str(pend_co.cell(r, 2).value).startswith("S2 ") and "NO liberar" in str(pend_co.cell(r, 4).value):
            bloqueos.append(["Cocina", "S2", "", "Soporte de campana y abatible"] + [""] * 11 + [pend_co.cell(r, 3).value])

    wb = openpyxl.Workbook()
    # Resumen
    ws = wb.active; ws.title = "Resumen"
    titulo(ws, "Compras del proyecto · cocina, clóset y baño",
           f"Estimación conjunta del {FECHA}, generada desde los Excel de cada mueble (herramientas/build_compras_proyecto.py).")
    encabezado(ws, 6, ["Concepto", "Por separado", "Conjunto", "Nota"], [42, 16, 16, 90])
    cocina_sep = libro("despiece-cocina.xlsx")["Compras"]["C2"].value
    closet_sep = libro("despiece-closet.xlsx")["Compras"]["C7"].value
    sep_rh = cocina_sep + closet_sep + 1
    resumen = [
        ["Láminas RH15", sep_rh, len(lam_rh), f"Cocina {cocina_sep} + clóset {closet_sep} + baño 1 por separado; encaje conjunto en la hoja Láminas."],
        ["Láminas de fondo RH6", 2, len(lam_fondo), "El fondo del baño cabe en la lámina de fondo de la cocina."],
        ["Rollos de canto rígido 100 m", 2, canto["2"]["rollos"], f"Más {canto['2']['sueltos']} m sueltos en la sucursal."],
        ["Rollos de canto flexible 200 m", 2, canto["0.45"]["rollos"], f"Sobran unos {canto['0.45']['sobrante']} m para cantear piezas pequeñas en el armado."],
        ["Bisagras HBM897-01 (pares)", 15, 15, "Misma referencia en cocina y baño: se compran juntas."],
    ]
    for i, f in enumerate(resumen, 7):
        for j, v in enumerate(f, 1): ws.cell(i, j, v)
    n = 7 + len(resumen) + 1
    ws.cell(n, 1, "Total con precio conocido (hoja Compras)").font = Font(bold=True)
    ws.cell(n, 3, "=Compras!G{}".format(10 + len(filas_compras) + len(grupos))).font = Font(bold=True)
    ws.cell(n, 3).number_format = '"$"#,##0'
    ws.cell(n, 4, "Los metros sueltos de rígido y las partidas sin precio no están incluidos.")
    ws.cell(n + 1, 1, "Partidas sin precio")
    ws.cell(n + 1, 3, "=Compras!G{}".format(11 + len(filas_compras) + len(grupos)))
    ws.cell(n + 3, 1, "Espejo, retal de granito y mesón son de otros proveedores; lavamanos, grifería y equipos se eligen aparte.").font = Font(italic=True)

    # Pedido de corte
    ws = wb.create_sheet("Pedido de corte")
    titulo(ws, "Pedido de corte y canteado · pedido único de los tres muebles",
           "Medidas finales en mm, con canto incluido; largo en sentido de veta. R = rígido 2 mm · F = flexible 0,45 mm · — = sin canto. Un solo tipo de canto por pieza, únicamente en los bordes indicados.")
    ws["A4"] = ("Supuestos de cada despiece: el proveedor pide medida terminada y descuenta el canto; piezas pequeñas = lado menor < 100 mm "
                "o mayor < 250 mm. Si la sucursal pide medida de corte, usar la columna «Corte sin canto».")
    encabezado(ws, 6, ["Mueble", "ID", "Módulo", "Pieza", "Material", "Cant.", "Largo a pedir", "Ancho a pedir", "Largo 1", "Largo 2",
                       "Corto 1", "Corto 2", "Canteo", "Medida terminada", "Corte sin canto", "Observaciones"],
               [9, 8, 18, 34, 22, 6, 10, 10, 7, 7, 7, 7, 10, 14, 15, 60])
    for i, f in enumerate(filas_pedido, 7):
        for j, v in enumerate(f, 1):
            c = ws.cell(i, j, v); c.alignment = Alignment(wrap_text=j in (4, 16), vertical="top")
    ws.freeze_panes = "A7"; ws.auto_filter.ref = f"A6:P{6 + len(filas_pedido)}"

    # Láminas
    ws = wb.create_sheet("Láminas")
    titulo(ws, "Encaje conjunto en láminas · propuesta de factibilidad",
           "Cortes pasantes, sin rotar por la veta, refilado 10 mm por lado y sierra 5 mm. El plano de corte definitivo lo define la seccionadora del proveedor.")
    encabezado(ws, 6, ["Material", "Lámina", "Ocupación", "Piezas"], [12, 8, 11, 120])
    r = 7
    for mat, lams in (("RH15", lam_rh), ("Fondo RH6", lam_fondo)):
        for k, lam in enumerate(lams, 1):
            ws.cell(r, 1, mat); ws.cell(r, 2, k); ws.cell(r, 3, round(ocupacion(lam), 3)).number_format = "0%"
            ws.cell(r, 4, ", ".join(p[0] for p in lam["items"])).alignment = Alignment(wrap_text=True, vertical="top"); r += 1

    # Cantos
    ws = wb.create_sheet("Cantos")
    titulo(ws, "Canto conjunto", f"Metros netos de cada despiece; compra con 5 cm por borde y {int(MARGEN*100)} % de margen.")
    encabezado(ws, 6, ["Tipo", "Mueble", "Metros netos", "Bordes"], [18, 22, 14, 10])
    r = 7
    for tipo, nombre in (("2", "Rígido 2 mm"), ("0.45", "Flexible 0,45 mm")):
        for m, neto, bordes in canto[tipo]["por_mueble"]:
            ws.cell(r, 1, nombre); ws.cell(r, 2, m); ws.cell(r, 3, round(neto, 3)); ws.cell(r, 4, bordes); r += 1
        c = canto[tipo]
        ws.cell(r, 1, nombre).font = Font(bold=True); ws.cell(r, 2, f"Total · comprar {c['necesario']} m").font = Font(bold=True)
        ws.cell(r, 3, round(c["neto"], 3)).font = Font(bold=True); ws.cell(r, 4, c["bordes"]).font = Font(bold=True); r += 2

    # Compras
    ws = wb.create_sheet("Compras")
    titulo(ws, "Compras consolidadas", "Precios de referencia de los Excel de cada mueble. «cotizar» = sin precio; no equivale a cero.")
    encabezado(ws, 9, ["Grupo", "Ítem", "Mueble", "Cant.", "Unidad", "Precio unitario (COP)", "Subtotal (COP)", "Nota"],
               [20, 46, 20, 7, 9, 14, 15, 80])
    r = 10
    for f in filas_compras:
        for j, v in enumerate(f[:6], 1): ws.cell(r, j, v)
        ws.cell(r, 7, f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(F{r})),D{r}*F{r},"Por cotizar")')
        ws.cell(r, 8, f[6]).alignment = Alignment(wrap_text=True, vertical="top")
        for c in (6, 7): ws.cell(r, c).number_format = '"$"#,##0'
        r += 1
    last = r - 1
    for g in grupos:
        ws.cell(r, 2, f"Subtotal {g}"); ws.cell(r, 7, f'=SUMIF(A10:A{last},"{g}",G10:G{last})').number_format = '"$"#,##0'; r += 1
    ws.cell(r, 2, "TOTAL CON PRECIO CONOCIDO").font = Font(bold=True)
    ws.cell(r, 7, f"=SUM(G{last + 1}:G{r - 1})").number_format = '"$"#,##0'; ws.cell(r, 7).font = Font(bold=True)
    ws.cell(r + 1, 2, "Partidas sin precio"); ws.cell(r + 1, 7, f'=COUNTIF(G10:G{last},"Por cotizar")')

    # Bloqueos
    ws = wb.create_sheet("Bloqueos")
    titulo(ws, "Antes de cortar", "Piezas que no deben cortarse hasta resolver lo indicado. Detalle en la hoja Pendientes de cada mueble.")
    encabezado(ws, 6, ["Mueble", "ID", "Pieza", "Qué falta"], [10, 10, 40, 100])
    for i, f in enumerate(bloqueos, 7):
        for j, v in enumerate((f[0], f[1], f[3], f[15]), 1): ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")

    # En el armado
    ws = wb.create_sheet("En el armado")
    titulo(ws, "Lo que se corta, recorta o rebaja en el armado",
           "Piezas que no se piden (se cortan de un sobrante), piezas pedidas con sobremedida para recortar y mecanizados en obra.")
    encabezado(ws, 6, ["Mueble", "ID", "Pieza", "Qué se hace"], [10, 10, 40, 110])
    for i, f in enumerate(armado, 7):
        for j, v in enumerate((f[0], f[1], f[3], f[15]), 1): ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(RAIZ / "compras-proyecto.xlsx")

    # Vista web
    total = sum(s for f in filas_compras if (s := subtotal(f)) is not None)
    sin_precio = sum(1 for f in filas_compras if subtotal(f) is None)
    b = ['<p><a href="./">← Proyecto</a> · <a href="compras-proyecto.xlsx" download>Descargar Excel</a></p>',
         '<span class="tag">Cocina · clóset · baño</span><h1>Compras del proyecto</h1>',
         '<nav class="tabs"><a href="#resumen">Resumen</a><a href="#pedido">Pedido de corte</a><a href="#laminas">Láminas</a>'
         '<a href="#cantos">Cantos</a><a href="#compras">Compras</a><a href="#bloqueos">Antes de cortar</a><a href="#armado">En el armado</a></nav>',
         '<h2 id="resumen">Resumen</h2>',
         tabla([r_[:4] for r_ in resumen], ["Concepto", "Por separado", "Conjunto", "Nota"], num_cols=(1, 2), wrap_cols=(3,)),
         f'<p class="note"><b>Total con precio conocido: {fmt(total, cop=True)}</b>; {sin_precio} partidas sin precio. '
         'Los metros sueltos de canto rígido se cotizan en la sucursal.</p>',
         '<h2 id="pedido">Pedido de corte</h2>',
         f'<p class="note">{html.escape(wb["Pedido de corte"]["A3"].value)} {html.escape(wb["Pedido de corte"]["A4"].value)}</p>',
         tabla(filas_pedido, [c.value for c in wb["Pedido de corte"][6]], num_cols=(5, 6, 7), wrap_cols=(15,)),
         '<h2 id="laminas">Láminas</h2>',
         f'<p class="note">{html.escape(wb["Láminas"]["A3"].value)}</p>',
         tabla([[m_, k, f"{ocupacion(l):.0%}", ", ".join(p[0] for p in l["items"])] for m_, ls in (("RH15", lam_rh), ("Fondo RH6", lam_fondo))
                for k, l in enumerate(ls, 1)], ["Material", "Lámina", "Ocupación", "Piezas"], wrap_cols=(3,)),
         '<h2 id="cantos">Cantos</h2>',
         tabla([[n_, m_, round(neto, 2), bordes] for t, n_ in (("2", "Rígido 2 mm"), ("0.45", "Flexible 0,45 mm"))
                for m_, neto, bordes in canto[t]["por_mueble"]] +
               [[f"Total {n_}", f"comprar {canto[t]['necesario']} m", round(canto[t]["neto"], 2), canto[t]["bordes"]]
                for t, n_ in (("2", "rígido"), ("0.45", "flexible"))],
               ["Tipo", "Mueble", "Metros netos", "Bordes"], num_cols=(2, 3), tot_pred=lambda r_: str(r_[0]).startswith("Total")),
         '<h2 id="compras">Compras</h2>',
         tabla([f[:6] + [subtotal(f) if subtotal(f) is not None else "Por cotizar", f[6]] for f in filas_compras],
               ["Grupo", "Ítem", "Mueble", "Cant.", "Unidad", "Precio unitario", "Subtotal", "Nota"],
               num_cols=(3,), cop_cols=(5, 6), wrap_cols=(7,)),
         '<h2 id="bloqueos">Antes de cortar</h2>',
         tabla([[f[0], f[1], f[3], f[15]] for f in bloqueos], ["Mueble", "ID", "Pieza", "Qué falta"], wrap_cols=(3,)),
         '<h2 id="armado">En el armado</h2>',
         f'<p class="note">{html.escape(wb["En el armado"]["A3"].value)}</p>',
         tabla([[f[0], f[1], f[3], f[15]] for f in armado], ["Mueble", "ID", "Pieza", "Qué se hace"], wrap_cols=(3,)),
         '<footer>Generado desde los Excel de cada mueble con <code>herramientas/build_compras_proyecto.py</code>. '
         'Vista de solo lectura.</footer>']
    (RAIZ / "compras-proyecto.html").write_text(page("Compras del proyecto", "".join(b)), encoding="utf-8", newline="\n")
    print(f"OK  compras-proyecto.xlsx y .html · RH15 {len(lam_rh)} láminas, fondo {len(lam_fondo)}, "
          f"rígido {canto['2']['necesario']} m, flexible {canto['0.45']['necesario']} m, total conocido {fmt(total, cop=True)}, "
          f"{sin_precio} sin precio, {len(filas_pedido)} líneas de pedido, {len(bloqueos)} bloqueos, {len(armado)} en el armado")


if __name__ == "__main__":
    main()
