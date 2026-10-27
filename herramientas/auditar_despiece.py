"""Auditoría automática de una versión de cocina: Excel de despiece  <->  modelo 3D.

Uso, desde la raíz del repositorio:

    python herramientas/auditar_despiece.py              # versión "propuesta"
    python herramientas/auditar_despiece.py propuesta 4  # versión y nº de láminas RH15 previstas

Requisitos: Python 3.9+ y openpyxl (pip install openpyxl). Solo lee archivos; no modifica nada.

Qué comprueba (cada bloque imprime OK / ATENCIÓN):
  1. Que despiece-<v>.xlsx y assets/models/cocina-<v>-data.js describan las MISMAS piezas
     (id, módulo, largo, ancho, cantidad, cantos, material).
  2. Que cada pieza 3D tenga las dimensiones de su fila del despiece y que haya tantas como "Cant.".
  3. Que dos piezas no ocupen el mismo volumen (choques de más de 0,5 mm).
  4. Que los metros de canto recalculados coincidan con la hoja Cantos.
  5. Que las medidas de cada módulo de la hoja Distribución sean consistentes con sus piezas
     (ancho del módulo = largo del piso + 30) y que los totales tecleados coincidan con el conteo.
  6. Que las piezas quepan en N láminas de 2440 x 2150 con cortes pasantes (guillotina),
     sin rotar (veta), refilado 10 mm por lado y sierra 4 mm. Es una prueba de factibilidad,
     no el plano de corte definitivo: la seccionadora del proveedor manda.

Convenciones del Excel (hoja Despiece): E = largo (veta), F = ancho, G = cantidad,
H/I = cantos de los lados largos, J/K = cantos de los lados cortos. "2" = rígido, "0.45" = flexible.
"""
import itertools, json, math, random, sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
slug = sys.argv[1] if len(sys.argv) > 1 else "propuesta"
LAMINAS = int(sys.argv[2]) if len(sys.argv) > 2 else 4
ESPESOR = {"Melamina Cartagena RH 15": 15, "Fondo melamina 6": 6}
errores = 0


def titulo(t):
    print(f"\n== {t}")


def ok(msg):
    print("  OK        " + msg)


def aviso(msg):
    global errores
    errores += 1
    print("  ATENCIÓN  " + msg)


# ---------------------------------------------------------------- carga
xlsx = RAIZ / f"despiece-{slug}.xlsx"
js = RAIZ / "assets" / "models" / f"cocina-{slug}-data.js"
if not xlsx.exists() or not js.exists():
    sys.exit(f"Faltan {xlsx.name} o {js.name}")
txt = js.read_text(encoding="utf-8")
datos = json.loads(txt[txt.index("{"): txt.rindex("}") + 1])
partes = {p["id"]: p for p in datos["parts"]}

wb = openpyxl.load_workbook(xlsx, data_only=True)  # valores que Excel guardó al grabar
filas = [r for r in wb["Despiece"].iter_rows(min_row=5, values_only=True) if r[0] and str(r[0]).startswith("P")]
if filas and filas[0][11] is None:
    sys.exit("El Excel no trae valores calculados: ábrelo en Excel, guárdalo y repite.")
xl = {r[0]: r for r in filas}

# ---------------------------------------------------------------- 1 y 2
titulo("1. Excel <-> modelo 3D")
solo_xl, solo_js = sorted(set(xl) - set(partes)), sorted(set(partes) - set(xl))
if solo_xl or solo_js:
    aviso(f"piezas solo en Excel {solo_xl}; solo en 3D {solo_js}")
dif = 0
for pid in sorted(set(xl) & set(partes)):
    r, p = xl[pid], partes[pid]
    pares = [("módulo", r[1].split()[0], p["module"]), ("largo", r[4], p["L"]), ("ancho", r[5], p["A"]),
             ("cantidad", r[6], p["qty"]), ("cantos", [str(x) for x in r[7:11]], p["edges"]),
             ("material", r[3], p["material"])]
    for n, a, b in pares:
        if a != b:
            dif += 1
            aviso(f"{pid} {n}: Excel={a} / 3D={b}")
    if len(p["placements"]) != p["qty"]:
        dif += 1
        aviso(f"{pid}: {len(p['placements'])} piezas en 3D, Cant.={p['qty']}")
    for a in p["placements"]:
        real = sorted(round(v, 1) for v in (a[1] - a[0], a[3] - a[2], a[5] - a[4]))
        esp = sorted([p["L"], p["A"], ESPESOR[p["material"]]])
        if any(abs(x - y) > 0.6 for x, y in zip(real, esp)):
            dif += 1
            aviso(f"{pid} medidas 3D {real} ≠ despiece {esp}")
if not dif:
    rh = sum(r[6] for r in filas if "RH" in r[3])
    fo = sum(r[6] for r in filas if "Fondo" in r[3])
    ok(f"{len(filas)} líneas, {rh} piezas RH15 + {fo} fondos: Excel y 3D coinciden")

# ---------------------------------------------------------------- 3
titulo("2. Choques entre piezas")
cajas = [(p["id"], i, p["name"], a) for p in datos["parts"] for i, a in enumerate(p["placements"])]
choques = 0
for A, B in itertools.combinations(cajas, 2):
    a, b = A[3], B[3]
    x = min(a[1], b[1]) - max(a[0], b[0])
    z = min(a[3], b[3]) - max(a[2], b[2])
    y = min(a[5], b[5]) - max(a[4], b[4])
    if x > 0.5 and z > 0.5 and y > 0.5:
        choques += 1
        aviso(f"{A[0]} {A[2][:30]} × {B[0]} {B[2][:30]}: solape {x:.1f} × {z:.1f} × {y:.1f} mm")
if not choques:
    ok(f"{len(cajas)} piezas, ningún volumen compartido")

# ---------------------------------------------------------------- 4
titulo("3. Cantos")
rig = fle = 0.0
for p in datos["parts"]:
    if "RH" not in p["material"]:
        continue
    for e, largo in zip(p["edges"], (p["L"], p["L"], p["A"], p["A"])):
        if e == "2":
            rig += largo * p["qty"] / 1000
        elif e == "0.45":
            fle += largo * p["qty"] / 1000
hoja = wb["Cantos"]
for nombre, calc, celda in (("rígido", rig, "C4"), ("flexible", fle, "C5")):
    v = hoja[celda].value
    if v is not None and abs(v - calc) < 0.01:
        ok(f"canto {nombre}: {calc:.3f} m")
    else:
        aviso(f"canto {nombre}: recalculado {calc:.3f} m, Excel {v}")

# ---------------------------------------------------------------- 5
titulo("4. Distribución y totales tecleados")
d = wb["Distribución"]
ancho_modulo = {}
for r in d.iter_rows(min_row=12, values_only=True):
    if r[0] and r[3]:
        ancho_modulo[str(r[0]).split()[0]] = r[3]
piso_base = {"B": ("Piso", 30), "S": ("Techo", 30)}
for mod, w in ancho_modulo.items():
    if mod[0] not in "BS" or mod == "B6":
        continue
    cand = [p for p in datos["parts"] if p["module"] == mod and p["name"].split()[0] == piso_base[mod[0]][0]]
    if cand and abs(cand[0]["L"] + 30 - w) > 0.01:
        aviso(f"{mod}: ancho en Distribución {w} ≠ {cand[0]['name']} {cand[0]['L']} + 30")
for etiqueta, valor in (("Piezas RH15", sum(r[6] for r in filas if "RH" in r[3])),
                        ("Fondos 6 mm", sum(r[6] for r in filas if "Fondo" in r[3]))):
    celda = next((r[1] for r in d.iter_rows(values_only=True) if r[0] == etiqueta), None)
    if celda == valor:
        ok(f"{etiqueta}: {valor}")
    else:
        aviso(f"{etiqueta}: Distribución dice {celda}, conteo real {valor}")
fe = openpyxl.load_workbook(xlsx, data_only=False)["Distribución"]
fijos = [c.coordinate for row in fe.iter_rows(min_row=25, max_row=29) for c in row[1:2]
         if c.value is not None and not str(c.value).startswith("=")]
if fijos:
    print(f"  NOTA      celdas tecleadas a mano (no se actualizan solas): {', '.join(fijos)}")

# ---------------------------------------------------------------- 6
titulo(f"5. Encaje en {LAMINAS} láminas RH15 (guillotina, sin rotar)")
K, BW, BH = 4, 2440 - 20, 2150 - 20


def corte(p):
    f = lambda e: 2 if e == "2" else 0.45 if e == "0.45" else 0
    return p["L"] - f(p["edges"][2]) - f(p["edges"][3]), p["A"] - f(p["edges"][0]) - f(p["edges"][1])


piezas = [(p["id"],) + corte(p) for p in datos["parts"] if "RH" in p["material"] for _ in range(p["qty"])]


def empacar(orden, elegir, dividir):
    laminas = []
    for pid, w, h in orden:
        W, H = w + K, h + K
        for lam in laminas + [None]:
            if lam is None:
                lam = {"libre": [(0, 0, BW + K, BH + K)], "items": []}
                laminas.append(lam)
            mejor = None
            for i, (x, y, fw, fh) in enumerate(lam["libre"]):
                if W <= fw and H <= fh:
                    s = {"baf": fw * fh - W * H, "bssf": min(fw - W, fh - H), "blsf": max(fw - W, fh - H)}[elegir]
                    if mejor is None or s < mejor[0]:
                        mejor = (s, i)
            if mejor:
                x, y, fw, fh = lam["libre"].pop(mejor[1])
                lam["items"].append((pid, x, y, w, h))
                rw, rh = fw - W, fh - H
                horiz = {"sas": rw < rh, "las": rw >= rh}[dividir]
                if horiz:
                    if rw > 0: lam["libre"].append((x + W, y, rw, H))
                    if rh > 0: lam["libre"].append((x, y + H, fw, rh))
                else:
                    if rw > 0: lam["libre"].append((x + W, y, rw, fh))
                    if rh > 0: lam["libre"].append((x, y + H, W, rh))
                break
            if lam is laminas[-1] and not lam["items"]:
                return None
    return laminas


rnd = random.Random(7)
ordenes = [sorted(piezas, key=lambda p: -p[1] * p[2]), sorted(piezas, key=lambda p: -p[1]), sorted(piezas, key=lambda p: -p[2])]
for _ in range(600):
    o = list(piezas)
    rnd.shuffle(o)
    ordenes.append(o)
mejor = None
for o in ordenes:
    for e in ("baf", "bssf", "blsf"):
        for s in ("sas", "las"):
            r = empacar(o, e, s)
            if r and (mejor is None or len(r) < len(mejor)):
                mejor = r
area_util = BW * BH
area_piezas = sum(w * h for _, w, h in piezas)
print(f"  área de piezas {area_piezas/1e6:.2f} m²  ({100*area_piezas/(LAMINAS*area_util):.0f}% de {LAMINAS} láminas)")
if mejor and len(mejor) <= LAMINAS:
    ok(f"caben en {len(mejor)} lámina(s); ocupación por lámina: "
       + ", ".join(f"{100*sum(w*h for _,_,_,w,h in l['items'])/area_util:.0f}%" for l in mejor))
else:
    aviso(f"no encontré encaje en {LAMINAS} láminas (mejor intento: {len(mejor) if mejor else '—'})")

print("\n" + ("RESULTADO: sin hallazgos." if not errores else f"RESULTADO: {errores} hallazgo(s) por revisar."))
sys.exit(1 if errores else 0)
