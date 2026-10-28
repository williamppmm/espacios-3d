"""Auditoría automática del despiece del clóset principal (solo lee el Excel; no modifica nada).

Uso, desde la raíz del repositorio:

    python herramientas/auditar_closet.py          # despiece-closet.xlsx
    python herramientas/auditar_closet.py 5        # y nº de láminas RH15 previstas

Requisitos: Python 3.9+ y openpyxl (pip install openpyxl).

Qué comprueba (cada bloque imprime OK / ATENCIÓN):
  1. Que cada fila de Despiece tenga cantos válidos ("2", "0.45" o "—") y que el área neta cuadre.
  2. Que los metros de canto recalculados coincidan con la hoja Cantos.
  3. Que las piezas quepan en N láminas de 2440 x 2150 (guillotina, sin rotar por la veta,
     refilado 10 mm por lado, sierra 4 mm). Es una prueba de factibilidad, no el plano de corte.
  4. Flecha estimada de las piezas anchas de RH15 (estantes, pisos de maletero, bases de cajón)
     con los refuerzos del despiece, frente al criterio L/200. Cálculo orientativo:
     apoyo simple, E = 2000 MPa (aglomerado RH, supuesto prudente), fluencia x2.
Los pesos de uso son supuestos editables al inicio del bloque 4.
El 3D no se lee aquí: compárelo abriendo closet.html (conteo de piezas e inventario).
"""
import math, random, sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
LAMINAS = int(sys.argv[1]) if len(sys.argv) > 1 else None
errores = 0


def titulo(t): print(f"\n== {t}")
def ok(m): print("  OK        " + m)
def aviso(m):
    global errores
    errores += 1
    print("  ATENCIÓN  " + m)


xlsx = RAIZ / "despiece-closet.xlsx"
if not xlsx.exists():
    sys.exit(f"Falta {xlsx.name}")
wb = openpyxl.load_workbook(xlsx, data_only=True)
ws = wb["Despiece"]
f = lambda e: 2 if e == "2" else 0.45 if e == "0.45" else 0
piezas = []   # (id, nombre, cant, L, A, cantos, esp)
for r in range(7, ws.max_row + 1):
    if not (ws.cell(r, 1).value and str(ws.cell(r, 1).value).startswith("C")):
        continue
    v = [ws.cell(r, c).value for c in range(1, 17)]
    piezas.append(dict(id=v[0], nombre=v[2], cant=v[3], L=v[4], A=v[5], esp=v[6], area=v[7], cantos=tuple(str(x) for x in v[8:12]), cl=v[12], ca=v[13]))

titulo("1. Filas del despiece")
malos = [p["id"] for p in piezas if any(c not in ("2", "0.45", "—") for c in p["cantos"])]
(aviso if malos else ok)(f"cantos inválidos en {malos}" if malos else f"{len(piezas)} líneas con cantos válidos")
n = sum(p["cant"] for p in piezas)
area = sum(p["cant"] * p["L"] * p["A"] for p in piezas) / 1e6
tot_a = sum(p["area"] * 1 for p in piezas)
print(f"  {n} piezas, {area:.6f} m² netos")
(ok if abs(sum(p['area'] for p in piezas) - area) < 1e-6 else aviso)("área de las filas = cantidad × largo × ancho")
errs = [p["id"] for p in piezas if abs(p["cl"] - (p["L"] - f(p["cantos"][2]) - f(p["cantos"][3]))) > 1e-6 or abs(p["ca"] - (p["A"] - f(p["cantos"][0]) - f(p["cantos"][1]))) > 1e-6]
(aviso if errs else ok)(f"corte teórico mal calculado en {errs}" if errs else "corte teórico = terminado − cantos (convención de la cocina)")
if any(p["esp"] != 15 for p in piezas): aviso("hay piezas que no son de 15 mm")

titulo("2. Metros de canto")
wc = wb["Cantos"]
res = float(wc["I2"].value)
neto = {"2": 0.0, "0.45": 0.0}
bordes = {"2": 0, "0.45": 0}
for p in piezas:
    for i, e in enumerate(p["cantos"]):
        if e in neto:
            neto[e] += p["cant"] * (p["L"] if i < 2 else p["A"]) / 1000
            bordes[e] += p["cant"]
for fila, e in ((7, "2"), (8, "0.45")):
    guardado = wc.cell(fila, 3).value
    pedir = math.ceil((neto[e] + bordes[e] * res) / 5) * 5
    (ok if abs(guardado - neto[e]) < 1e-6 and wc.cell(fila, 6).value == pedir else aviso)(
        f"canto {'rígido' if e == '2' else 'flexible'}: {neto[e]:.3f} m netos, {bordes[e]} bordes, pedir {pedir} m (Excel: {guardado:.3f} / {wc.cell(fila, 6).value})")

titulo("3. Encaje en láminas RH15 (guillotina, sin rotar)")
K, BW, BH = 4, 2440 - 20, 2150 - 20
lista = [(p["id"], p["cl"], p["ca"]) for p in piezas for _ in range(p["cant"])]


def empacar(orden, elegir, dividir):
    laminas = []
    for pid, w, h in orden:
        W, H = w + K, h + K
        for lam in laminas + [None]:
            if lam is None:
                lam = {"libre": [(0, 0, BW + K, BH + K)], "items": []}
                laminas.append(lam)
            mejor = None
            for i, (x, y, fw, fh) in enumerate(lam["libre"]):
                if W <= fw and H <= fh:
                    s = {"baf": fw * fh - W * H, "bssf": min(fw - W, fh - H), "blsf": max(fw - W, fh - H)}[elegir]
                    if mejor is None or s < mejor[0]:
                        mejor = (s, i)
            if mejor:
                x, y, fw, fh = lam["libre"].pop(mejor[1])
                lam["items"].append((pid, x, y, w, h))
                rw, rh = fw - W, fh - H
                horiz = {"sas": rw < rh, "las": rw >= rh}[dividir]
                if horiz:
                    if rw > 0: lam["libre"].append((x + W, y, rw, H))
                    if rh > 0: lam["libre"].append((x, y + H, fw, rh))
                else:
                    if rw > 0: lam["libre"].append((x + W, y, rw, fh))
                    if rh > 0: lam["libre"].append((x, y + H, W, rh))
                break
            if lam is laminas[-1] and not lam["items"]:
                return None
    return laminas


rnd = random.Random(7)
ordenes = [sorted(lista, key=lambda p: -p[1] * p[2]), sorted(lista, key=lambda p: -p[1]), sorted(lista, key=lambda p: -p[2])]
for _ in range(600):
    o = list(lista); rnd.shuffle(o); ordenes.append(o)
mejor = None
for o in ordenes:
    for e in ("baf", "bssf", "blsf"):
        for s in ("sas", "las"):
            r = empacar(o, e, s)
            if r and (mejor is None or len(r) < len(mejor)):
                mejor = r
au = BW * BH
if mejor is None:
    aviso("alguna pieza no cabe en una lámina")
else:
    print(f"  mejor encaje encontrado: {len(mejor)} láminas; ocupación: " + ", ".join(f"{100*sum(w*h for _,_,_,w,h in l['items'])/au:.0f}%" for l in mejor))
    if LAMINAS is None: ok("encaje calculado (indique el nº previsto como 2.º argumento para compararlo)")
    elif len(mejor) <= LAMINAS: ok(f"caben en las {LAMINAS} láminas previstas")
    else: aviso(f"no encontré encaje en {LAMINAS} láminas (mejor: {len(mejor)})")

titulo("4. Estimaciones teóricas de flexión (L/200, E=2000 MPa, fluencia x2)")
print("  Hipótesis: reparto de carga y unión rígida ideal. No verifica tornillos, adhesivo, anclajes ni compatibilidad de herrajes.")
g, dens, E = 9.81, 650e-9, 2000.0


def flecha(L, kg, I): return 2 * 5 * (kg * g / L) * L**4 / (384 * E * I)


def T(b, h, bw, hw):
    A1, A2 = b * h, bw * hw
    y1, y2 = hw + h / 2, hw / 2
    yc = (A1 * y1 + A2 * y2) / (A1 + A2)
    return b * h**3 / 12 + A1 * (y1 - yc) ** 2 + bw * hw**3 / 12 + A2 * (y2 - yc) ** 2


casos = [
    ("Piso de maletero 990×480 + listón frontal 80 (viga de borde, mitad de carga), 20 kg", 990,
     lambda: flecha(990, (20 + 990*480*15*dens + 990*(80+100)*15*dens) / 2, 15 * 80**3 / 12)),
    ("Repisa 990×348 + 2 listones de 60, 10 kg", 990,
     lambda: flecha(990, (10 + 990*348*15*dens + 2*990*60*15*dens) / 2, 15 * 60**3 / 12)),
    ("Base de cajón 793×400 + refuerzo central 15×30, 15 kg", 763,
     lambda: flecha(763, 15 + 763*400*15*dens + 763*30*15*dens, T(400, 15, 15, 30))),
]
for nombre, L, fn in casos:
    d = fn()
    (ok if d <= L / 200 else aviso)(f"{nombre}: {d:.1f} mm (límite {L/200:.1f} mm)")

print("\nRESULTADO NUMÉRICO: " + ("comprobaciones correctas." if not errores else f"{errores} hallazgo(s) por revisar."))
print("FABRICACIÓN: pendiente de confirmar herrajes, uniones, anclajes y acabado. Una comprobación numérica correcta no libera el pedido de corte.")
sys.exit(1 if errores else 0)
